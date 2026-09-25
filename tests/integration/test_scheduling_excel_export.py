"""Testes de integração do endpoint de export Excel do Agendamento (HML-242).

Valida GET /scheduling/export/excel: content-type/headers, espelhamento da tela
(toggle de terminais + filtro de status da legenda) e preenchimento das colunas
que dependem da O.S. finalizada.
"""

from datetime import UTC, date, datetime
from io import BytesIO

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.employees.models import Employee
from app.modules.scheduling.models import Appointment
from app.modules.service_orders.enums import OSStatus
from app.modules.service_orders.models import (
    ServiceOrder,
    ServiceOrderItem,
    ServiceOrderWorker,
)
from app.modules.services.models import Service
from app.modules.stores.models import Store

pytestmark = pytest.mark.asyncio

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest_asyncio.fixture
async def film_service(db_session: AsyncSession, test_brand) -> Service:
    svc = Service(
        name="Película Lateral",
        code="PL01",
        department="film",
        base_price=300.00,
        is_active=True,
        brand_id=test_brand.id,
    )
    db_session.add(svc)
    await db_session.commit()
    await db_session.refresh(svc)
    return svc


async def _add_appointment(
    db: AsyncSession,
    store: Store,
    owner_id: int,
    plate: str,
    *,
    delivery: date,
    service_order_id: int | None = None,
) -> Appointment:
    appt = Appointment(
        store_id=store.id,
        department="film",
        delivery_date=delivery,
        vehicle_plate=plate,
        vehicle_model="Corolla",
        vehicle_color="Branco",
        service_ids=None,
        film_entries=[{"service_id": 1, "tonality": "G20"}],
        status="scheduled",
        created_by_id=owner_id,
        service_order_id=service_order_id,
    )
    db.add(appt)
    await db.commit()
    await db.refresh(appt)
    return appt


def _load(content: bytes):
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    return header, rows


class TestExcelExportEndpoint:
    async def test_returns_xlsx_with_headers(
        self, owner_client: AsyncClient, db_session, test_store, test_owner
    ):
        await _add_appointment(
            db_session, test_store, test_owner.id, "AAA1D23", delivery=date(2030, 6, 15)
        )
        resp = await owner_client.get("/api/v1/scheduling/export/excel")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(XLSX_MIME)
        assert "attachment" in resp.headers["content-disposition"]
        assert ".xlsx" in resp.headers["content-disposition"]

        header, rows = _load(resp.content)
        assert header[0] == "Status Agendamento"
        assert len(header) == 21
        plates = {r[5] for r in rows}
        assert "AAA1D23" in plates

    async def test_hides_completed_without_include_terminal(
        self,
        owner_client: AsyncClient,
        db_session: AsyncSession,
        test_store,
        test_owner,
        film_service,
    ):
        """Espelha a tela: sem include_terminal, O.S. finalizada não aparece."""
        so = ServiceOrder(
            store_id=test_store.id,
            vehicle_plate="DONE123",
            department="film",
            status=OSStatus.COMPLETED.value,
            created_by_id=test_owner.id,
            updated_by_id=test_owner.id,
            photos="[]",
            entry_time=datetime.now(UTC),
        )
        db_session.add(so)
        await db_session.flush()
        db_session.add(
            ServiceOrderItem(
                service_order_id=so.id,
                service_id=film_service.id,
                quantity=1,
                unit_price=300.00,
            )
        )
        await db_session.commit()
        await _add_appointment(
            db_session,
            test_store,
            test_owner.id,
            "DONE123",
            delivery=date(2030, 6, 15),
            service_order_id=so.id,
        )
        await _add_appointment(
            db_session, test_store, test_owner.id, "PEND123", delivery=date(2030, 6, 15)
        )

        # Sem include_terminal → só o pendente
        resp = await owner_client.get("/api/v1/scheduling/export/excel")
        _, rows = _load(resp.content)
        plates = {r[5] for r in rows}
        assert "PEND123" in plates
        assert "DONE123" not in plates

        # Com include_terminal → aparece o finalizado
        resp2 = await owner_client.get(
            "/api/v1/scheduling/export/excel", params={"include_terminal": True}
        )
        _, rows2 = _load(resp2.content)
        plates2 = {r[5] for r in rows2}
        assert "DONE123" in plates2

    async def test_display_status_filter_mirrors_legend(
        self, owner_client: AsyncClient, db_session, test_store, test_owner
    ):
        """display_status restringe às linhas do status selecionado na legenda."""
        # Atrasado (data passada, sem O.S.)
        await _add_appointment(
            db_session, test_store, test_owner.id, "LATE123", delivery=date(2020, 1, 1)
        )
        # Agendado (data futura, sem O.S.)
        await _add_appointment(
            db_session, test_store, test_owner.id, "FUT1D23", delivery=date(2030, 6, 15)
        )
        resp = await owner_client.get(
            "/api/v1/scheduling/export/excel", params={"display_status": ["atrasado"]}
        )
        _, rows = _load(resp.content)
        plates = {r[5] for r in rows}
        assert "LATE123" in plates
        assert "FUT1D23" not in plates

    async def test_completed_os_fills_value_and_installer(
        self,
        owner_client: AsyncClient,
        db_session: AsyncSession,
        test_store,
        test_owner,
        film_service,
    ):
        """O.S. finalizada preenche Valor (soma) e Instalador (workers)."""
        so = ServiceOrder(
            store_id=test_store.id,
            vehicle_plate="VAL1D23",
            department="film",
            status=OSStatus.COMPLETED.value,
            created_by_id=test_owner.id,
            updated_by_id=test_owner.id,
            photos="[]",
            entry_time=datetime.now(UTC),
        )
        db_session.add(so)
        await db_session.flush()
        db_session.add(
            ServiceOrderItem(
                service_order_id=so.id,
                service_id=film_service.id,
                quantity=2,
                unit_price=300.00,
            )
        )
        emp = Employee(name="Carlos Instalador", store_id=test_store.id, is_active=True)
        db_session.add(emp)
        await db_session.flush()
        db_session.add(ServiceOrderWorker(service_order_id=so.id, employee_id=emp.id))
        await db_session.commit()

        await _add_appointment(
            db_session,
            test_store,
            test_owner.id,
            "VAL1D23",
            delivery=date(2030, 6, 15),
            service_order_id=so.id,
        )
        resp = await owner_client.get(
            "/api/v1/scheduling/export/excel", params={"include_terminal": True}
        )
        header, rows = _load(resp.content)
        row = next(r for r in rows if r[5] == "VAL1D23")
        assert row[16] == 600.0  # Valor = 2 × 300
        assert row[15] == "Carlos Instalador"  # Instalador
