"""Testes unitários do export Excel do Agendamento (HML-242).

Cobre a função pura ``build_scheduling_excel_matrix`` (mapeamento das 21 colunas
do modelo RAGS) e o gerador ``generate_scheduling_excel`` (openpyxl → bytes).
"""

from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.modules.scheduling.excel import (
    EXCEL_COLUMNS,
    build_scheduling_excel_matrix,
    generate_scheduling_excel,
)
from app.modules.scheduling.schemas import AppointmentResponse


def _make_appt(**overrides) -> AppointmentResponse:
    """AppointmentResponse mínimo para os testes, sobrescrevendo o que interessa."""
    base = {
        "id": 1,
        "store_id": 10,
        "store_name": "BYD Unidade 08",
        "department": "film",
        "delivery_date": date(2026, 8, 18),
        "delivery_time": time(14, 0, 0),
        "external_os_number": "7265",
        "vehicle_plate": "V0034888",
        "vehicle_model": "Dolphin",
        "vehicle_color": "Preto",
        "consultant_id": None,
        "consultant_name": "Roberta De Moraes Areias",
        "service_ids": [1, 2],
        "service_names": ["WP 1 - Pelicula Lateral", "WP 2 - Pelicula Parabrisa"],
        "notes": "ENTREGA DIA 18/08",
        "is_galpon": False,
        "is_courtesy": False,
        "is_return": False,
        "film_type_id": None,
        "film_tonality": None,
        "film_entries": None,
        "status": "scheduled",
        "display_status": "em_execucao",
        "service_order_id": 555,
        "service_order_number": "BYD-2608-0007",
        "created_by_id": 99,
        "cancelled_at": None,
        "cancellation_reason": None,
        "created_at": datetime(2026, 8, 14, 15, 12, tzinfo=UTC),  # 12:12 BRT
        "updated_at": datetime(2026, 8, 17, 14, 44, tzinfo=UTC),  # 11:44 BRT
    }
    base.update(overrides)
    return AppointmentResponse.model_validate(base)


class TestExcelColumns:
    def test_has_21_columns_in_model_order(self):
        assert EXCEL_COLUMNS == [
            "Status Agendamento",
            "Data Agendamento",
            "Loja",
            "Local",
            "DPTO",
            "Placa/ Chassi",
            "O.S",
            "Cortesia?",
            "Retorno?",
            "Consultor",
            "Modelo V.",
            "Cor",
            "Observações",
            "Serviços",
            "Cód Rolo",
            "Instalador",
            "Valor",
            "Data de Cadastro",
            "Responsável Cadastro",
            "Data Alteração",
            "Responsável Alteração",
        ]


class TestBuildMatrix:
    def test_maps_appointment_fields(self):
        appt = _make_appt()
        matrix = build_scheduling_excel_matrix(
            [appt],
            os_extras={},
            user_names={99: "Luana Gonçalves"},
            last_changes={1: (datetime(2026, 8, 17, 14, 44, tzinfo=UTC), 42)},
        )
        assert len(matrix) == 1
        row = matrix[0]
        assert row[0] == "Em execução"  # Status Agendamento
        assert row[1] == "18/08/2026 às 14:00:00"  # Data Agendamento
        assert row[2] == "BYD Unidade 08"  # Loja
        assert row[3] == ""  # Local (não galpão)
        assert row[4] == "Película"  # DPTO
        assert row[5] == "V0034888"  # Placa/Chassi
        assert row[6] == "7265"  # O.S = external_os_number
        assert row[7] == "Não"  # Cortesia?
        assert row[8] == "Não"  # Retorno?
        assert row[9] == "Roberta De Moraes Areias"  # Consultor
        assert row[10] == "Dolphin"  # Modelo V.
        assert row[11] == "Preto"  # Cor
        assert row[12] == "ENTREGA DIA 18/08"  # Observações
        assert row[13] == "WP 1 - Pelicula Lateral\nWP 2 - Pelicula Parabrisa"  # Serviços
        assert row[17] == "14/08/2026, 12:12"  # Data de Cadastro (BRT)
        assert row[18] == "Luana Gonçalves"  # Responsável Cadastro

    def test_local_shows_galpao_when_is_galpon(self):
        appt = _make_appt(is_galpon=True)
        row = build_scheduling_excel_matrix([appt], {}, {}, {})[0]
        assert row[3] == "Galpão"

    def test_courtesy_and_return_flags(self):
        appt = _make_appt(is_courtesy=True, is_return=True)
        row = build_scheduling_excel_matrix([appt], {}, {}, {})[0]
        assert row[7] == "Sim"  # Cortesia?
        assert row[8] == "Sim"  # Retorno?

    def test_os_dependent_columns_empty_without_extras(self):
        """Sem O.S. finalizada: Cód Rolo, Instalador e Valor ficam vazios."""
        appt = _make_appt()
        row = build_scheduling_excel_matrix([appt], os_extras={}, user_names={}, last_changes={})[0]
        assert row[14] == ""  # Cód Rolo
        assert row[15] == ""  # Instalador
        assert row[16] == ""  # Valor

    def test_os_dependent_columns_filled_when_completed(self):
        """O.S. finalizada popula Valor (soma) e Instalador; Cód Rolo vem das film_entries."""
        appt = _make_appt(
            display_status="finalizado",
            film_entries=[{"service_id": 1, "tonality": "G20", "film_roll_code": "BYD-CE-2608-01"}],
        )
        os_extras = {555: {"value": Decimal("930.00"), "installers": ["João", "Pedro"]}}
        row = build_scheduling_excel_matrix([appt], os_extras, {}, {})[0]
        assert row[14] == "BYD-CE-2608-01"  # Cód Rolo
        assert row[15] == "João, Pedro"  # Instalador
        assert row[16] == 930.0  # Valor (numérico)

    def test_last_change_columns(self):
        appt = _make_appt()
        last = {1: (datetime(2026, 8, 17, 14, 44, tzinfo=UTC), 42)}
        row = build_scheduling_excel_matrix([appt], {}, {42: "Angelo Marcio"}, last)[0]
        assert row[19] == "17/08/2026, 11:44"  # Data Alteração (BRT)
        assert row[20] == "Angelo Marcio"  # Responsável Alteração

    def test_no_last_change_leaves_columns_empty(self):
        appt = _make_appt()
        row = build_scheduling_excel_matrix([appt], {}, {}, {})[0]
        assert row[19] == ""  # Data Alteração
        assert row[20] == ""  # Responsável Alteração

    def test_delivery_without_time(self):
        appt = _make_appt(delivery_time=None)
        row = build_scheduling_excel_matrix([appt], {}, {}, {})[0]
        assert row[1] == "18/08/2026"


class TestGenerateExcel:
    def test_produces_valid_xlsx_with_header(self):
        matrix = [["Em execução"] + [""] * 20]
        data = generate_scheduling_excel(matrix)
        assert isinstance(data, bytes) and len(data) > 0
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert header == EXCEL_COLUMNS
        assert ws.cell(row=2, column=1).value == "Em execução"

    def test_empty_matrix_still_has_header(self):
        data = generate_scheduling_excel([])
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        assert [c.value for c in ws[1]] == EXCEL_COLUMNS
        assert ws.max_row == 1
