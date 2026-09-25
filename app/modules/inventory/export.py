"""
Export Excel de saídas avulsas de película (workbook programático, sem template).
"""

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.modules.inventory.models import FilmWithdrawal
from app.modules.inventory.service import build_withdrawal_response_dict

TZ_LOCAL = ZoneInfo("America/Sao_Paulo")

_WITHDRAWAL_HEADERS = [
    "Data",
    "Loja",
    "Funcionário",
    "Tipo de Película",
    "Tonalidade",
    "Bobina",
    "Metros",
    "Motivo",
    "Registrado por",
    "Status",
    "Estornado por",
    "Estornado em",
]

_SUMMARY_HEADERS = ["Funcionário", "Qtde Saídas", "Total Metros"]


def _fmt_dt(dt) -> str:
    if dt is None:
        return ""
    if dt.tzinfo is not None:
        dt = dt.astimezone(TZ_LOCAL)
    return dt.strftime("%d/%m/%Y %H:%M")


def _fmt_roll_label(data: dict) -> str:
    """Bobina em formato legível: "06/04/2026 · 15m" (tipo/tonalidade já têm coluna)."""
    receipt = data.get("roll_receipt_date")
    total = data.get("roll_total_meters")
    if receipt is None:
        return data.get("roll_visual_id", "")
    label = receipt.strftime("%d/%m/%Y")
    if total is not None:
        label += f" · {total:g}m"
    return label


def generate_withdrawals_excel(withdrawals: list[FilmWithdrawal], summary: list[dict]) -> bytes:
    """
    Gera Excel com 2 abas: "Saídas" (linha a linha, inclui estornadas) e
    "Resumo por Funcionário" (exclui estornadas — base do desconto em folha).
    """
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Saídas"
    ws.append(_WITHDRAWAL_HEADERS)
    for cell in ws[1]:
        cell.font = bold
    ws.freeze_panes = "A2"

    for w in withdrawals:
        data = build_withdrawal_response_dict(w)
        ws.append(
            [
                _fmt_dt(data["created_at"]),
                data["store_name"] or "",
                data["employee_name"] or "",
                data["film_type_name"] or "",
                data["tonality"] or "",
                _fmt_roll_label(data),
                round(data["meters"], 2),
                data["reason"] or "",
                data["created_by_name"] or "",
                "Estornada" if data["is_reversed"] else "Ativa",
                data["reversed_by_name"] or "",
                _fmt_dt(data["reversed_at"]),
            ]
        )

    widths = [17, 20, 28, 22, 12, 32, 9, 40, 24, 11, 24, 17]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws_summary = wb.create_sheet("Resumo por Funcionário")
    ws_summary.append(_SUMMARY_HEADERS)
    for cell in ws_summary[1]:
        cell.font = bold
    ws_summary.freeze_panes = "A2"

    total_meters = 0.0
    for item in summary:
        total_meters += item["total_meters"]
        ws_summary.append(
            [item["employee_name"], item["withdrawal_count"], round(item["total_meters"], 2)]
        )

    total_row = ws_summary.max_row + 1
    ws_summary.cell(row=total_row, column=1, value="TOTAL").font = bold
    ws_summary.cell(
        row=total_row, column=2, value=sum(i["withdrawal_count"] for i in summary)
    ).font = bold
    ws_summary.cell(row=total_row, column=3, value=round(total_meters, 2)).font = bold

    for idx, width in enumerate([28, 12, 14], start=1):
        ws_summary.column_dimensions[get_column_letter(idx)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
