"""Export Excel do Agendamento (HML-242).

Reproduz a planilha operacional "RAGS": uma linha por agendamento exibido na
tela, com 21 colunas (status, data, loja, veículo, serviços, O.S., valor,
instalador e rastro de cadastro/alteração).

As colunas que dependem da finalização da O.S. (Cód Rolo, Instalador, Valor)
ficam vazias enquanto o trabalho não foi concluído — igual ao comportamento da
planilha original.
"""

from datetime import datetime, time
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.scheduling.schemas import AppointmentResponse

TZ_LOCAL = ZoneInfo("America/Sao_Paulo")

# Cabeçalho na ordem exata do modelo (21 colunas).
EXCEL_COLUMNS = [
    "Status Agendamento",
    "Data Agendamento",
    "Loja",
    "Local",
    "DPTO",
    "Placa/ Chassi",
    "O.S",
    "Cortesia?",
    "Retorno?",
    "Consultor",
    "Modelo V.",
    "Cor",
    "Observações",
    "Serviços",
    "Cód Rolo",
    "Instalador",
    "Valor",
    "Data de Cadastro",
    "Responsável Cadastro",
    "Data Alteração",
    "Responsável Alteração",
]

STATUS_LABELS = {
    "atrasado": "Atrasado",
    "atencao": "Atenção",
    "agendado": "Agendado",
    "em_execucao": "Em execução",
    "finalizado": "Finalizado",
    "cancelado": "Cancelado",
    "duplicidade": "Duplicidade",
}

DEPARTMENT_LABELS = {
    "film": "Película",
    "security_film": "Película de Segurança",
    "ppf": "PPF",
    "bodywork": "Funilaria",
    "vn": "VN",
    "vd": "VD",
    "vu": "VU",
    "workshop": "Oficina",
}

# Colunas com quebra de linha (índice 0-based): Observações e Serviços.
_WRAP_COLUMNS = {12, 13}

# Larguras aproximadas do modelo original.
_COLUMN_WIDTHS = [
    24,
    22,
    16,
    10,
    12,
    17,
    9,
    10,
    10,
    24,
    15,
    9,
    28,
    44,
    15,
    16,
    11,
    18,
    22,
    18,
    22,
]

_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _to_local(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    return dt.astimezone(TZ_LOCAL) if dt.tzinfo else dt


def _delivery_label(delivery_date, delivery_time: time | None) -> str:
    """Data de entrega: "18/08/2026 às 14:00:00" (ou só a data sem horário)."""
    day = delivery_date.strftime("%d/%m/%Y")
    if delivery_time is None:
        return day
    return f"{day} às {delivery_time.strftime('%H:%M:%S')}"


def _timestamp_label(dt: datetime | None) -> str:
    """Carimbo de cadastro/alteração no fuso local: "14/08/2026, 12:12"."""
    local = _to_local(dt)
    return local.strftime("%d/%m/%Y, %H:%M") if local else ""


def _roll_codes(appt: AppointmentResponse) -> str:
    """Códigos das bobinas usadas (das film_entries já enriquecidas), sem repetir.

    Vazios até o Finalizar atribuir a bobina — self-gating para O.S. não concluída.
    """
    codes: list[str] = []
    for fe in appt.film_entries or []:
        applications = fe.get("applications") or []
        if applications:
            for app in applications:
                code = app.get("film_roll_code")
                if code and code not in codes:
                    codes.append(code)
        code = fe.get("film_roll_code")
        if code and code not in codes:
            codes.append(code)
    return ", ".join(codes)


def build_scheduling_excel_matrix(
    appointments: list[AppointmentResponse],
    os_extras: dict[int, dict],
    user_names: dict[int, str],
    last_changes: dict[int, tuple[datetime, int | None]],
) -> list[list]:
    """Monta a matriz (linha por agendamento × 21 colunas) na ordem do modelo.

    Args:
        appointments: agendamentos já convertidos (com film_entries enriquecidas).
        os_extras: {service_order_id: {"value": Decimal|None, "installers": [nomes]}}
            — preenchido apenas para O.S. finalizadas (Valor/Instalador).
        user_names: {user_id: nome} para Responsável Cadastro/Alteração.
        last_changes: {appointment_id: (data_alteração, user_id)} da última edição.
    """
    matrix: list[list] = []
    for appt in appointments:
        extras = os_extras.get(appt.service_order_id) if appt.service_order_id else None
        installers = extras.get("installers") if extras else None
        value = extras.get("value") if extras else None

        change = last_changes.get(appt.id)
        change_date = _timestamp_label(change[0]) if change else ""
        change_user = user_names.get(change[1], "") if change and change[1] else ""

        matrix.append(
            [
                STATUS_LABELS.get(appt.display_status, appt.display_status),
                _delivery_label(appt.delivery_date, appt.delivery_time),
                appt.store_name or "",
                "Galpão" if appt.is_galpon else "",
                DEPARTMENT_LABELS.get(appt.department, appt.department),
                appt.vehicle_plate or "",
                appt.external_os_number or "",
                "Sim" if appt.is_courtesy else "Não",
                "Sim" if appt.is_return else "Não",
                appt.consultant_name or "",
                appt.vehicle_model or "",
                appt.vehicle_color or "",
                appt.notes or "",
                "\n".join(appt.service_names) if appt.service_names else "",
                _roll_codes(appt),
                ", ".join(installers) if installers else "",
                float(value) if value is not None else "",
                _timestamp_label(appt.created_at),
                user_names.get(appt.created_by_id, "") if appt.created_by_id else "",
                change_date,
                change_user,
            ]
        )
    return matrix


def generate_scheduling_excel(matrix: list[list]) -> bytes:
    """Gera o workbook (aba "Agendamentos") com cabeçalho + linhas da matriz."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Agendamentos"

    for col, title in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(matrix, start=2):
        for c, value in enumerate(row):
            cell = ws.cell(row=r, column=c + 1, value=value)
            if c in _WRAP_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

    for i, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
