"""
Service Order export - Generates Excel reports from new branded templates.
"""

from collections.abc import Callable
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl

from app.modules.service_orders.enums import OSStatus
from app.modules.service_orders.models import ServiceOrder

_TEMPLATES = Path(__file__).parent.parent.parent.parent / "templates"

CONFERENCE_TEMPLATE_PATH = _TEMPLATES / "Relatório Conferência.xlsx"
GERAL_TEMPLATE_PATH = _TEMPLATES / "Relatório Fechamento Geral.xlsx"
VN_TEMPLATE_PATH = _TEMPLATES / "Relatório Fechamento VN.xlsx"
VU_TEMPLATE_PATH = _TEMPLATES / "Relatório Fechamento VU.xlsx"
FUNI_TEMPLATE_PATH = _TEMPLATES / "Relatório Fechamento Funilaria.xlsx"
OFICINA_TEMPLATE_PATH = _TEMPLATES / "Relatório Fechamento Oficina.xlsx"
RESUMO_TEMPLATE_PATH = _TEMPLATES / "Resumo do Fechamento Mensal.xlsx"
INVENTORY_TEMPLATE_PATH = _TEMPLATES / "Relatório de Estoque de Película.xlsx"
ROLL_TEMPLATE_PATH = _TEMPLATES / "Relatório de Película por Bobina.xlsx"

DEPT_LABELS: dict[str, str] = {
    "film": "Película",
    "security_film": "Película de Segurança",
    "ppf": "PPF",
    "vn": "VN",
    "vd": "Venda Direta",
    "vu": "VU",
    "bodywork": "Funilaria",
    "workshop": "Oficina",
}

_DEPTS_NO_EXT_OS = {"vn", "vd", "vu"}
_DEPTS_12COL = {"vn", "vu"}

# Status terminais prevalecem sobre is_verified na coluna "Status Conf."
_TERMINAL_STATUS_LABELS: dict[str, str] = {
    OSStatus.DUPLICATE.value: "Duplicado",
    OSStatus.CANCELLED.value: "Cancelada",
    OSStatus.WRONG.value: "Lançado Errado",
}

_DEPT_TEMPLATE_MAP: dict[str, Path] = {
    "vn": VN_TEMPLATE_PATH,
    "vu": VU_TEMPLATE_PATH,
    "bodywork": FUNI_TEMPLATE_PATH,
    "workshop": OFICINA_TEMPLATE_PATH,
}


def _reset_view(wb: openpyxl.Workbook) -> None:
    for sheet in wb.worksheets:
        sheet.sheet_view.view = "normal"


def _format_date(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    try:
        return value.strftime("%d/%m/%Y")
    except AttributeError:
        return str(value)


def _period_str(date_from: datetime | None, date_to: datetime | None) -> str:
    if date_from and date_to:
        f, t = date_from.strftime("%m/%Y"), date_to.strftime("%m/%Y")
        return f if f == t else f"{f} a {t}"
    if date_from:
        return date_from.strftime("%m/%Y")
    if date_to:
        return date_to.strftime("%m/%Y")
    return ""


def _clean_notes(raw: str) -> str:
    return (
        raw.replace("[CORTESIA]", "")
        .replace("[RETORNO]", "")
        .replace(" | ", " ")
        .strip(" |")
        .strip()
    )


def _write_cell(ws, row: int, col: int, value, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if fmt:
        cell.number_format = fmt


def _update_table_ref(ws, last_data_row: int, last_col_letter: str, start_row: int = 1) -> None:
    ref = f"A{start_row}:{last_col_letter}{last_data_row}"
    for tbl in ws.tables.values():
        tbl.ref = ref
        if tbl.autoFilter is not None:
            tbl.autoFilter.ref = ref


# ── Conferência ──────────────────────────────────────────────────────────────


def generate_conference_excel(orders: list[ServiceOrder]) -> bytes:
    """
    Relatório Conferência — 27 colunas (A–AA).

    A: Data da O.S  B: Loja         C: Local        D: DPTO
    E: Placa/Chassi F: O.S          G: Cortesia?    H: Retorno?
    I: Consultor    J: Modelo Veículo K: Cor         L: Observações
    M: NFe          N: Cód Serviço  O: Desc Serviço P: Tonalidade
    Q: Cód Rolo     R: Instalador   S: Valor        T: ✔
    U: Obs Conferência V: Status Conf. W: Status Veículo
    X: Data Cadastro Y: Responsável Cadastro
    Z: Data Alteração AA: Responsável Alteração
    """
    wb = openpyxl.load_workbook(CONFERENCE_TEMPLATE_PATH)
    _reset_view(wb)
    ws = wb.active

    start_row = 2
    row_idx = start_row

    for order in orders:
        store_nm = order.store.name if order.store else ""
        local = "Galpão" if order.is_galpon else ""
        service_date = order.service_date or (order.entry_time.date() if order.entry_time else None)
        date_str = _format_date(service_date)
        dept_label = DEPT_LABELS.get(order.department or "", order.department or "")
        dept_key = (order.department or "").lower()

        if order.external_os_number:
            os_number = order.external_os_number
        elif dept_key in _DEPTS_NO_EXT_OS:
            os_number = ""
        else:
            os_number = order.order_number or f"#{order.id}"

        plate = order.vehicle_plate or ""
        is_courtesy = "Sim" if order.is_courtesy else "Não"
        is_return = "Sim" if order.is_return else "Não"
        consultant_name = (
            order.consultant.name if order.consultant else (order.consultant_name or "")
        )
        vehicle_full = f"{order.vehicle_brand or ''} {order.vehicle_model or ''}".strip()
        color = order.vehicle_color or ""
        notes = _clean_notes(order.notes or "")
        invoice = order.invoice_number or ""
        is_verified_mark = "✓" if order.is_verified else ""
        status_conf = _TERMINAL_STATUS_LABELS.get(
            order.status, "Verificado" if order.is_verified else "Pendente"
        )
        worker_names = ", ".join(w.employee.name for w in (order.workers or []) if w.employee)
        created_at_str = _format_date(order.created_at)
        updated_at_str = _format_date(order.updated_at)
        created_by_name = order.created_by.full_name if order.created_by else ""

        items = order.items or []
        if items:
            for item in items:
                svc = item.service
                item_value = float(item.unit_price or 0) * int(item.quantity or 1)
                _write_conference_row(
                    ws,
                    row_idx,
                    date_str,
                    store_nm,
                    local,
                    dept_label,
                    plate,
                    os_number,
                    is_courtesy,
                    is_return,
                    consultant_name,
                    vehicle_full,
                    color,
                    notes,
                    invoice,
                    (svc.code or "") if svc else "",
                    (svc.name or "") if svc else "",
                    item.tonality or "",
                    item.roll_code or "",
                    worker_names,
                    item_value,
                    is_verified_mark,
                    status_conf,
                    created_at_str,
                    created_by_name,
                    updated_at_str,
                )
                row_idx += 1
        else:
            _write_conference_row(
                ws,
                row_idx,
                date_str,
                store_nm,
                local,
                dept_label,
                plate,
                os_number,
                is_courtesy,
                is_return,
                consultant_name,
                vehicle_full,
                color,
                notes,
                invoice,
                "",
                "",
                "",
                "",
                worker_names,
                0.0,
                is_verified_mark,
                status_conf,
                created_at_str,
                created_by_name,
                updated_at_str,
            )
            row_idx += 1

    last_data_row = row_idx - 1
    if last_data_row >= start_row:
        _update_table_ref(ws, last_data_row, "AA")
        subtotal_row = last_data_row + 1
        sc = ws.cell(row=subtotal_row, column=19)  # S: Valor
        sc.value = f"=SUBTOTAL(9,S{start_row}:S{last_data_row})"
        sc.number_format = "#,##0.00"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _write_conference_row(
    ws,
    row_idx: int,
    date_str: str,
    store_name: str,
    local: str,
    dept_label: str,
    plate: str,
    os_number: str,
    is_courtesy: str,
    is_return: str,
    consultant_name: str,
    vehicle_full: str,
    color: str,
    notes: str,
    invoice: str,
    service_code: str,
    service_desc: str,
    tonality: str,
    roll_code: str,
    worker_names: str,
    item_value: float,
    is_verified_mark: str,
    status_conf: str,
    created_at_str: str,
    created_by_name: str,
    updated_at_str: str,
) -> None:
    ws.cell(row=row_idx, column=1).value = date_str  # A: Data da O.S
    ws.cell(row=row_idx, column=2).value = store_name  # B: Loja
    ws.cell(row=row_idx, column=3).value = local  # C: Local
    ws.cell(row=row_idx, column=4).value = dept_label  # D: DPTO
    ws.cell(row=row_idx, column=5).value = plate  # E: Placa/Chassi
    ws.cell(row=row_idx, column=6).value = os_number  # F: O.S
    ws.cell(row=row_idx, column=7).value = is_courtesy  # G: Cortesia?
    ws.cell(row=row_idx, column=8).value = is_return  # H: Retorno?
    ws.cell(row=row_idx, column=9).value = consultant_name  # I: Consultor
    ws.cell(row=row_idx, column=10).value = vehicle_full  # J: Modelo Veículo
    ws.cell(row=row_idx, column=11).value = color  # K: Cor
    ws.cell(row=row_idx, column=12).value = notes  # L: Observações
    ws.cell(row=row_idx, column=13).value = invoice  # M: NFe
    ws.cell(row=row_idx, column=14).value = service_code  # N: Cód Serviço
    ws.cell(row=row_idx, column=15).value = service_desc  # O: Descrição Serviço
    ws.cell(row=row_idx, column=16).value = tonality  # P: Tonalidade
    ws.cell(row=row_idx, column=17).value = roll_code  # Q: Cód Rolo
    ws.cell(row=row_idx, column=18).value = worker_names  # R: Instalador
    vc = ws.cell(row=row_idx, column=19)  # S: Valor
    vc.value = item_value
    vc.number_format = "#,##0.00"
    ws.cell(row=row_idx, column=20).value = is_verified_mark  # T: ✔
    ws.cell(row=row_idx, column=21).value = ""  # U: Obs Conferência
    ws.cell(row=row_idx, column=22).value = status_conf  # V: Status Conf.
    ws.cell(row=row_idx, column=23).value = ""  # W: Status Veículo
    ws.cell(row=row_idx, column=24).value = created_at_str  # X: Data Cadastro
    ws.cell(row=row_idx, column=25).value = created_by_name  # Y: Responsável Cadastro
    ws.cell(row=row_idx, column=26).value = updated_at_str  # Z: Data Alteração
    ws.cell(row=row_idx, column=27).value = ""  # AA: Responsável Alteração


# ── Fechamento tabular ────────────────────────────────────────────────────────


def generate_fechamento_excel(
    orders: list[ServiceOrder],
    store_name: str,
    date_from: datetime | None,
    date_to: datetime | None,
    department: str | None = None,
    item_filter: Callable[..., bool] | None = None,
) -> bytes:
    """
    Relatório tabular de fechamento com seleção automática de template por departamento.

    vn/vu   → 12 colunas (A–L, sem coluna O.S.)
    demais  → 13 colunas (A–M, com coluna O.S.)

    item_filter: predicado opcional chamado por item; quando fornecido, apenas os itens
    que passam são incluídos na linha do Excel. O.S. sem itens (já sem nenhum item
    que case) produz uma linha zerada (comportamento original).
    """
    template_path = _DEPT_TEMPLATE_MAP.get(department or "", GERAL_TEMPLATE_PATH)
    wb = openpyxl.load_workbook(template_path)
    _reset_view(wb)
    ws = wb.active

    use_12col = department in _DEPTS_12COL
    start_row = 2
    row_idx = start_row

    for order in orders:
        service_date = order.service_date or (order.entry_time.date() if order.entry_time else None)
        date_str = _format_date(service_date)
        dept_label = DEPT_LABELS.get(order.department or "", order.department or "")
        store_nm = order.store.name if order.store else store_name
        local = "Galpão" if order.is_galpon else ""
        plate = order.vehicle_plate or ""
        color = order.vehicle_color or ""
        notes = _clean_notes(order.notes or "")
        consultant_name = (
            order.consultant.name if order.consultant else (order.consultant_name or "")
        )
        vehicle_full = f"{order.vehicle_brand or ''} {order.vehicle_model or ''}".strip()

        dept_key = (order.department or "").lower()
        if order.external_os_number:
            os_number = order.external_os_number
        elif dept_key in _DEPTS_NO_EXT_OS:
            os_number = ""
        else:
            os_number = order.order_number or f"#{order.id}"

        all_items = order.items or []
        # Apply item-level filter without mutating the ORM relationship
        items = [i for i in all_items if item_filter(i)] if item_filter is not None else all_items

        if items:
            for item in items:
                svc = item.service
                item_value = float(item.unit_price or 0) * int(item.quantity or 1)
                if use_12col:
                    _write_fechamento_row_12(
                        ws,
                        row_idx,
                        date_str,
                        store_nm,
                        local,
                        dept_label,
                        plate,
                        consultant_name,
                        vehicle_full,
                        color,
                        notes,
                        (svc.code or "") if svc else "",
                        (svc.name or "") if svc else "",
                        item_value,
                    )
                else:
                    _write_fechamento_row_13(
                        ws,
                        row_idx,
                        date_str,
                        store_nm,
                        local,
                        dept_label,
                        plate,
                        os_number,
                        consultant_name,
                        vehicle_full,
                        color,
                        notes,
                        (svc.code or "") if svc else "",
                        (svc.name or "") if svc else "",
                        item_value,
                    )
                row_idx += 1
        else:
            if use_12col:
                _write_fechamento_row_12(
                    ws,
                    row_idx,
                    date_str,
                    store_nm,
                    local,
                    dept_label,
                    plate,
                    consultant_name,
                    vehicle_full,
                    color,
                    notes,
                    "",
                    "",
                    0.0,
                )
            else:
                _write_fechamento_row_13(
                    ws,
                    row_idx,
                    date_str,
                    store_nm,
                    local,
                    dept_label,
                    plate,
                    os_number,
                    consultant_name,
                    vehicle_full,
                    color,
                    notes,
                    "",
                    "",
                    0.0,
                )
            row_idx += 1

    last_data_row = row_idx - 1
    last_col_letter = "L" if use_12col else "M"
    value_col = 12 if use_12col else 13
    count_col_letter = "J" if use_12col else "K"

    if last_data_row >= start_row:
        _update_table_ref(ws, last_data_row, last_col_letter)
        subtotal_row = last_data_row + 1
        ws.cell(
            row=subtotal_row, column=value_col - 2
        ).value = f"=SUBTOTAL(3,{count_col_letter}{start_row}:{count_col_letter}{last_data_row})"
        svc = ws.cell(row=subtotal_row, column=value_col)
        svc.value = f"=SUBTOTAL(9,{last_col_letter}{start_row}:{last_col_letter}{last_data_row})"
        svc.number_format = "#,##0.00"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _write_fechamento_row_13(
    ws,
    row_idx: int,
    date_str: str,
    store_name: str,
    local: str,
    dept_label: str,
    plate: str,
    os_number: str,
    consultant_name: str,
    vehicle_full: str,
    color: str,
    notes: str,
    service_code: str,
    service_desc: str,
    item_value: float,
) -> None:
    ws.cell(row=row_idx, column=1).value = date_str  # A: Data da O.S
    ws.cell(row=row_idx, column=2).value = store_name  # B: Loja
    ws.cell(row=row_idx, column=3).value = local  # C: Local
    ws.cell(row=row_idx, column=4).value = dept_label  # D: DPTO
    ws.cell(row=row_idx, column=5).value = plate  # E: Placa/Chassi
    ws.cell(row=row_idx, column=6).value = os_number  # F: O.S
    ws.cell(row=row_idx, column=7).value = consultant_name  # G: Consultor
    ws.cell(row=row_idx, column=8).value = vehicle_full  # H: Modelo Veículo
    ws.cell(row=row_idx, column=9).value = color  # I: Cor
    ws.cell(row=row_idx, column=10).value = notes  # J: Observações
    ws.cell(row=row_idx, column=11).value = service_code  # K: Cód Serviço
    ws.cell(row=row_idx, column=12).value = service_desc  # L: Descrição Serviço
    vc = ws.cell(row=row_idx, column=13)  # M: Valor
    vc.value = item_value
    vc.number_format = "#,##0.00"


def _write_fechamento_row_12(
    ws,
    row_idx: int,
    date_str: str,
    store_name: str,
    local: str,
    dept_label: str,
    plate: str,
    consultant_name: str,
    vehicle_full: str,
    color: str,
    notes: str,
    service_code: str,
    service_desc: str,
    item_value: float,
) -> None:
    ws.cell(row=row_idx, column=1).value = date_str  # A: Data da O.S
    ws.cell(row=row_idx, column=2).value = store_name  # B: Loja
    ws.cell(row=row_idx, column=3).value = local  # C: Local
    ws.cell(row=row_idx, column=4).value = dept_label  # D: DPTO
    ws.cell(row=row_idx, column=5).value = plate  # E: Placa/Chassi
    ws.cell(row=row_idx, column=6).value = consultant_name  # F: Consultor
    ws.cell(row=row_idx, column=7).value = vehicle_full  # G: Modelo Veículo
    ws.cell(row=row_idx, column=8).value = color  # H: Cor
    ws.cell(row=row_idx, column=9).value = notes  # I: Observações
    ws.cell(row=row_idx, column=10).value = service_code  # J: Cód Serviço
    ws.cell(row=row_idx, column=11).value = service_desc  # K: Descrição Serviço
    vc = ws.cell(row=row_idx, column=12)  # L: Valor
    vc.value = item_value
    vc.number_format = "#,##0.00"


# ── Resumo (slot-based) ───────────────────────────────────────────────────────

_CORTESIA_SLOTS: list[tuple[int, int, int, int, int]] = [
    # (name_row, name_col, data_row, qnt_col, total_col)
    (9, 1, 10, 1, 2),
    (9, 3, 10, 3, 4),
    (9, 5, 10, 5, 6),
    (9, 7, 10, 7, 9),
    (9, 10, 10, 10, 11),
    (11, 1, 12, 1, 2),
    (11, 3, 12, 3, 4),
    (11, 5, 12, 5, 6),
    (11, 7, 12, 7, 9),
    (11, 10, 12, 10, 11),
    (13, 1, 14, 1, 2),
    (13, 3, 14, 3, 4),
    (13, 5, 14, 5, 6),
    (13, 7, 14, 7, 9),
    (13, 10, 14, 10, 11),
]

_FUNI_SLOTS: list[tuple] = [(19, 1, 20, 1, 2), (19, 3, 20, 3, 4), (19, 5, 20, 5, 6)]

_VN_SLOTS: list[tuple] = [
    (26, 1, 27, 1, 2),
    (26, 3, 27, 3, 4),
    (26, 5, 27, 5, 6),
    (28, 1, 29, 1, 2),
    (28, 3, 29, 3, 4),
    (28, 5, 29, 5, 6),
]

_VD_SLOTS: list[tuple] = [
    (33, 1, 34, 1, 2),
    (33, 3, 34, 3, 4),
    (33, 5, 34, 5, 6),
    (33, 7, 34, 7, 9),
    (33, 10, 34, 10, 11),
]

_VU_SLOTS: list[tuple] = [
    (40, 1, 41, 1, 2),
    (40, 3, 41, 3, 4),
    (40, 5, 41, 5, 6),
    (42, 1, 43, 1, 2),
    (42, 3, 43, 3, 4),
    (42, 5, 43, 5, 6),
]

_SECTION_ROWS: dict[str, tuple[int, int]] = {
    "workshop_courtesy": (2, 16),
    "workshop_lavagem": (2, 16),
    "bodywork": (17, 22),
    "vn": (23, 31),
    "vd": (32, 36),
    "vu": (37, 45),
}

_LAV_SIM_P = (
    "lavagem simples",
    "lavagem + aspiração",
    "lavagem com aspiração",
    "lav c/aspira",
    "lav. simples",
)
_LAV_COM_P = ("lavagem completa", "lavagem + aspiração + motor")
_LAV_TEST_P = ("test drive", "lav test", "lavagem test")
_DUCHA_P = ("ducha",)


def _matches(name: str, patterns: tuple) -> bool:
    nl = name.lower()
    return any(p in nl for p in patterns)


def _svc_agg(orders: list[ServiceOrder], dept_key: str) -> dict[str, dict]:
    """
    Agrega serviços por nome para um grupo de departamento.

    Para grupos de Oficina (workshop_courtesy, workshop_lavagem, workshop_other):
    - workshop_courtesy: agrega TODOS os itens das O.S. marcadas como cortesia
      (precedência D3 — O.S. cortesia inteira fica neste grupo).
    - workshop_lavagem: agrega apenas itens NÃO-cortesia cujo service.name contém
      "lavagem simples" (classificação por item, D1).
    - workshop_other: agrega apenas itens NÃO-cortesia cujo service.name NÃO contém
      "lavagem simples" (classificação por item, D1).

    Para demais departamentos, agrega todos os itens cuja O.S. pertence ao dept_key.
    """
    result: dict[str, dict] = {}

    def _add(svc_name: str, svc_code: str, value: float) -> None:
        if svc_name not in result:
            result[svc_name] = {"count": 0, "total": 0.0, "code": svc_code}
        result[svc_name]["count"] += 1
        result[svc_name]["total"] += value

    for order in orders:
        dept = (order.department or "").lower()
        items = order.items or []

        if dept_key in ("workshop_courtesy", "workshop_lavagem", "workshop_other"):
            if dept != "workshop":
                continue

            if dept_key == "workshop_courtesy":
                # D3: O.S. cortesia inteira → agrega todos os itens aqui
                if not order.is_courtesy:
                    continue
                for item in items:
                    svc_name = (item.service.name if item.service else None) or "Serviço sem nome"
                    svc_code = (item.service.code if item.service else None) or ""
                    value = float(item.unit_price or 0) * int(item.quantity or 1)
                    _add(svc_name, svc_code, value)
            else:
                # workshop_lavagem / workshop_other — classificação por item (D1)
                if order.is_courtesy:
                    continue
                for item in items:
                    is_lav = item.service and "lavagem simples" in (item.service.name or "").lower()
                    if dept_key == "workshop_lavagem" and not is_lav:
                        continue
                    if dept_key == "workshop_other" and is_lav:
                        continue
                    svc_name = (item.service.name if item.service else None) or "Serviço sem nome"
                    svc_code = (item.service.code if item.service else None) or ""
                    value = float(item.unit_price or 0) * int(item.quantity or 1)
                    _add(svc_name, svc_code, value)
        else:
            if dept != dept_key:
                continue
            for item in items:
                svc_name = (item.service.name if item.service else None) or "Serviço sem nome"
                svc_code = (item.service.code if item.service else None) or ""
                value = float(item.unit_price or 0) * int(item.quantity or 1)
                _add(svc_name, svc_code, value)

    return result


def _is_lav_funi(svc_name: str, svc_code: str) -> bool:
    name_l, code_l = svc_name.lower(), svc_code.lower()
    return "lav.funi" in code_l or ("lavagem" in name_l and "funi" in name_l)


def _freeze_formulas(ws, row_start: int, row_end: int) -> None:
    for r in range(row_start, row_end + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = 0


def _fill_oficina(ws, orders: list[ServiceOrder], extra_cortesias: dict | None = None) -> None:
    ws_lav = _svc_agg(orders, "workshop_lavagem")
    ws_cort = _svc_agg(orders, "workshop_courtesy")
    ws_other = _svc_agg(orders, "workshop_other")

    # H6: contar LAV.CORTESIA de todos os grupos (cortesia, lavagem e outros).
    # Com a classificação por item, itens cujo código é LAV.CORTESIA mas cujo nome
    # NÃO contém "lavagem simples" migram para workshop_other — somamos os três grupos.
    lav_count = 0
    ws_cort_display: dict[str, dict] = {}

    for k, v in ws_cort.items():
        if v.get("code", "").upper() == "LAV.CORTESIA":
            lav_count += v["count"]
        else:
            ws_cort_display[k] = v

    for _k, v in ws_lav.items():
        if v.get("code", "").upper() == "LAV.CORTESIA":
            lav_count += v["count"]

    for _k, v in ws_other.items():
        if v.get("code", "").upper() == "LAV.CORTESIA":
            lav_count += v["count"]

    # Adicionar serviços extras de funilaria nas cortesias
    if extra_cortesias:
        for k, v in extra_cortesias.items():
            if k in ws_cort_display:
                ws_cort_display[k]["count"] += v["count"]
                ws_cort_display[k]["total"] += v["total"]
            else:
                ws_cort_display[k] = dict(v)

    _write_cell(ws, 6, 1, None)
    _write_cell(ws, 6, 2, None)
    # D6 preservado: fórmula =B6*0.5 do template

    if lav_count > 0:
        _write_cell(ws, 6, 8, lav_count)
    # I6 preservado: fórmula =H6*30 do template

    for i, (svc_name, svc_data) in enumerate(list(ws_cort_display.items())[:15]):
        nr, nc, dr, qc, tc = _CORTESIA_SLOTS[i]
        code = svc_data.get("code", "")
        display = code if svc_name.count(" ") > 2 else svc_name
        _write_cell(ws, nr, nc, display)
        _write_cell(ws, dr, qc, svc_data["count"])
        _write_cell(ws, dr, tc, svc_data["total"], "#,##0.00")
    # C15 preservado: fórmula =SUM(B14,D14,...) do template


def _fill_funilaria(ws, orders: list[ServiceOrder]) -> dict:
    ws_funi = _svc_agg(orders, "bodywork")

    lav_funi: dict[str, dict] = {}
    extras: dict[str, dict] = {}

    for k, v in ws_funi.items():
        if _is_lav_funi(k, v.get("code", "")):
            lav_funi[k] = v
        else:
            extras[k] = v

    if lav_funi:
        nr, nc, dr, qc, _ = _FUNI_SLOTS[0]
        total_count = sum(v["count"] for v in lav_funi.values())
        svc_name = next(iter(lav_funi.keys()))
        _write_cell(ws, nr, nc, svc_name)
        _write_cell(ws, dr, qc, total_count)
        # B20 preservado: fórmula =A20*60 do template
    # C21 preservado: fórmula =SUM(F20,D20,B20) do template
    # J19 preservado: fórmula =SUM(C21) do template

    return extras


def _fill_vn(ws, orders: list[ServiceOrder]) -> None:
    ws_vn = _svc_agg(orders, "vn")
    vn_lav_sim = {k: v for k, v in ws_vn.items() if _matches(k, _LAV_SIM_P)}
    vn_lav_com = {k: v for k, v in ws_vn.items() if _matches(k, _LAV_COM_P)}
    vn_lav_test = {k: v for k, v in ws_vn.items() if _matches(k, _LAV_TEST_P)}
    vn_ducha = {k: v for k, v in ws_vn.items() if _matches(k, _DUCHA_P)}
    vn_reg = {
        k: v
        for k, v in ws_vn.items()
        if k not in vn_lav_sim
        and k not in vn_lav_com
        and k not in vn_lav_test
        and k not in vn_ducha
    }

    sim_val = sum(v["total"] for v in vn_lav_sim.values())
    com_val = sum(v["total"] for v in vn_lav_com.values())
    test_val = sum(v["total"] for v in vn_lav_test.values())
    dch_val = sum(v["total"] for v in vn_ducha.values())

    if vn_lav_sim:
        _write_cell(ws, 27, 8, sum(v["count"] for v in vn_lav_sim.values()))
        _write_cell(ws, 27, 9, sim_val, "#,##0.00")
    if vn_lav_com:
        _write_cell(ws, 27, 10, sum(v["count"] for v in vn_lav_com.values()))
        _write_cell(ws, 27, 11, com_val, "#,##0.00")
    if vn_lav_test:
        _write_cell(ws, 29, 8, sum(v["count"] for v in vn_lav_test.values()))
        _write_cell(ws, 29, 9, test_val, "#,##0.00")
    if vn_ducha:
        _write_cell(ws, 29, 10, sum(v["count"] for v in vn_ducha.values()))
        _write_cell(ws, 29, 11, dch_val, "#,##0.00")

    reg_total = 0.0
    for i, (svc_name, svc_data) in enumerate(list(vn_reg.items())[:6]):
        nr, nc, dr, qc, tc = _VN_SLOTS[i]
        _write_cell(ws, nr, nc, svc_name)
        _write_cell(ws, dr, qc, svc_data["count"])
        _write_cell(ws, dr, tc, svc_data["total"], "#,##0.00")
        reg_total += svc_data["total"]

    _write_cell(ws, 30, 3, reg_total, "#,##0.00")
    _write_cell(ws, 30, 10, sim_val + com_val + test_val + dch_val, "#,##0.00")


def _fill_vd(ws, orders: list[ServiceOrder]) -> None:
    ws_vd = _svc_agg(orders, "vd")
    total_val = 0.0
    for i, (svc_name, svc_data) in enumerate(list(ws_vd.items())[:5]):
        nr, nc, dr, qc, tc = _VD_SLOTS[i]
        _write_cell(ws, nr, nc, svc_name)
        _write_cell(ws, dr, qc, svc_data["count"])
        _write_cell(ws, dr, tc, svc_data["total"], "#,##0.00")
        total_val += svc_data["total"]
    _write_cell(ws, 35, 9, total_val, "#,##0.00")


def _fill_vu(ws, orders: list[ServiceOrder]) -> None:
    ws_vu = _svc_agg(orders, "vu")
    vu_lav_sim = {k: v for k, v in ws_vu.items() if _matches(k, _LAV_SIM_P)}
    vu_lav_com = {k: v for k, v in ws_vu.items() if _matches(k, _LAV_COM_P)}
    vu_ducha = {k: v for k, v in ws_vu.items() if _matches(k, _DUCHA_P)}
    vu_reg = {
        k: v
        for k, v in ws_vu.items()
        if k not in vu_lav_sim and k not in vu_lav_com and k not in vu_ducha
    }

    sim_val = sum(v["total"] for v in vu_lav_sim.values())
    com_val = sum(v["total"] for v in vu_lav_com.values())
    dch_val = sum(v["total"] for v in vu_ducha.values())

    if vu_lav_sim:
        _write_cell(ws, 41, 8, sum(v["count"] for v in vu_lav_sim.values()))
        _write_cell(ws, 41, 9, sim_val, "#,##0.00")
    if vu_lav_com:
        _write_cell(ws, 41, 10, sum(v["count"] for v in vu_lav_com.values()))
        _write_cell(ws, 41, 11, com_val, "#,##0.00")
    if vu_ducha:
        _write_cell(ws, 43, 8, sum(v["count"] for v in vu_ducha.values()))
        _write_cell(ws, 43, 9, dch_val, "#,##0.00")

    reg_total = 0.0
    for i, (svc_name, svc_data) in enumerate(list(vu_reg.items())[:6]):
        nr, nc, dr, qc, tc = _VU_SLOTS[i]
        _write_cell(ws, nr, nc, svc_name)
        _write_cell(ws, dr, qc, svc_data["count"])
        _write_cell(ws, dr, tc, svc_data["total"], "#,##0.00")
        reg_total += svc_data["total"]

    _write_cell(ws, 44, 3, reg_total, "#,##0.00")
    _write_cell(ws, 44, 10, sim_val + com_val + dch_val, "#,##0.00")


_SECTION_FILL: dict[str, Callable[..., Any]] = {
    "workshop_courtesy": _fill_oficina,
    "workshop_lavagem": _fill_oficina,
    "bodywork": _fill_funilaria,
    "vn": _fill_vn,
    "vd": _fill_vd,
    "vu": _fill_vu,
}


def _set_resumo_title(
    ws, store_name: str, date_from: datetime | None, date_to: datetime | None
) -> None:
    """Substitui os placeholders de loja e período na célula de título (linha 1)."""
    period = _period_str(date_from, date_to)
    title_cell = ws.cell(row=1, column=1)
    if title_cell.value:
        title_cell.value = (
            str(title_cell.value).replace("LOJA", store_name).replace("ABR - 2026", period)
        )


def generate_resumo_fechamento_excel(
    orders: list[ServiceOrder],
    store_name: str,
    date_from: datetime | None,
    date_to: datetime | None,
) -> bytes:
    wb = openpyxl.load_workbook(RESUMO_TEMPLATE_PATH)
    _reset_view(wb)
    ws = wb.active

    _set_resumo_title(ws, store_name, date_from, date_to)
    extras = _fill_funilaria(ws, orders)
    _fill_oficina(ws, orders, extra_cortesias=extras)
    _fill_vn(ws, orders)
    _fill_vd(ws, orders)
    _fill_vu(ws, orders)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_section_fechamento_excel(
    orders: list[ServiceOrder],
    dept_key: str,
    store_name: str,
    date_from: datetime | None,
    date_to: datetime | None,
) -> bytes:
    wb = openpyxl.load_workbook(RESUMO_TEMPLATE_PATH)
    _reset_view(wb)
    ws = wb.active

    _set_resumo_title(ws, store_name, date_from, date_to)

    fill_fn = _SECTION_FILL.get(dept_key)
    if fill_fn:
        fill_fn(ws, orders)

    section = _SECTION_ROWS.get(dept_key)
    if section:
        sec_start, sec_end = section
        max_row = max(ws.max_row, 56)
        for r in range(2, sec_start):
            ws.row_dimensions[r].hidden = True
        for r in range(sec_end + 1, max_row + 1):
            ws.row_dimensions[r].hidden = True

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Estoque de Película ───────────────────────────────────────────────────────


def generate_inventory_excel(rolls: list) -> bytes:
    """
    Relatório de Estoque de Película — 17 colunas (A–Q).

    A: Loja             B: Origem (vazio)   C: Tipo de Película D: Ton
    E: Data_Entrada      F: ID_Bobina        G: Fornecedor       H: NFE
    I: Valor (vazio)    J: Metro_Inicial    K: Metro_Atual      L: Status
    M: Serviços Feitos  N: Data Cadastro    O: Responsável Cadastro (vazio)
    P: Data Movimentação Q: Responsável Movimentação (vazio)
    """
    wb = openpyxl.load_workbook(INVENTORY_TEMPLATE_PATH)
    _reset_view(wb)
    ws = wb.active

    start_row = 2
    row_idx = start_row

    for roll in rolls:
        store_nm = roll.store.name if roll.store else ""
        film_type_nm = roll.film_type.name if roll.film_type else ""
        tonality = roll.tonality or ""
        receipt_date_str = _format_date(roll.receipt_date)
        visual_id = getattr(roll, "visual_id", None) or f"#{roll.id}"
        supplier = roll.supplier or ""
        nfe = roll.nfe_number or ""
        total_meters = float(roll.total_meters or 0)
        remaining_meters = float(roll.remaining_meters or 0)
        status = roll.status or ""
        consumptions_count = len(roll.consumptions) if roll.consumptions is not None else 0
        created_at_str = _format_date(roll.created_at)
        updated_at_str = _format_date(roll.updated_at)

        ws.cell(row=row_idx, column=1).value = store_nm  # A: Loja
        ws.cell(row=row_idx, column=2).value = ""  # B: Origem
        ws.cell(row=row_idx, column=3).value = film_type_nm  # C: Tipo de Película
        ws.cell(row=row_idx, column=4).value = tonality  # D: Ton
        ws.cell(row=row_idx, column=5).value = receipt_date_str  # E: Data_Entrada
        ws.cell(row=row_idx, column=6).value = visual_id  # F: ID_Bobina
        ws.cell(row=row_idx, column=7).value = supplier  # G: Fornecedor
        ws.cell(row=row_idx, column=8).value = nfe  # H: NFE
        ws.cell(row=row_idx, column=9).value = ""  # I: Valor (não rastreado)
        ws.cell(row=row_idx, column=10).value = total_meters  # J: Metro_Inicial
        ws.cell(row=row_idx, column=11).value = remaining_meters  # K: Metro_Atual
        ws.cell(row=row_idx, column=12).value = status  # L: Status
        ws.cell(row=row_idx, column=13).value = consumptions_count  # M: Serviços Feitos
        ws.cell(row=row_idx, column=14).value = created_at_str  # N: Data Cadastro
        ws.cell(row=row_idx, column=15).value = ""  # O: Responsável Cadastro
        ws.cell(row=row_idx, column=16).value = updated_at_str  # P: Data Movimentação
        ws.cell(row=row_idx, column=17).value = ""  # Q: Responsável Movimentação
        row_idx += 1

    last_data_row = row_idx - 1
    if last_data_row >= start_row:
        _update_table_ref(ws, last_data_row, "Q")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Película por Bobina ───────────────────────────────────────────────────────


def generate_roll_excel(roll, items_with_orders: list) -> bytes:
    """
    Relatório de Película por Bobina — 26 colunas (A–Z).

    A: Data da O.S  B: Loja         C: Local        D: DPTO
    E: Placa/Chassi F: O.S          G: Cortesia?    H: Retorno?
    I: Consultor    J: Modelo Veículo K: Cor         L: Observações
    M: NFe          N: Cód Serviço  O: Desc Serviço P: Tonalidade
    Q: Cód Bobina   R: ID_Bobina    S: Instalador   T: Valor
    U: ✔            V: Obs Conferência W: Data Cadastro
    X: Responsável Cadastro Y: Data Alteração Z: Responsável Alteração
    """
    wb = openpyxl.load_workbook(ROLL_TEMPLATE_PATH)
    _reset_view(wb)
    ws = wb.active

    roll_visual_id = getattr(roll, "visual_id", None) or f"#{roll.id}"

    start_row = 2
    row_idx = start_row

    for item in items_with_orders:
        order = item.service_order
        if order is None:
            continue

        service_date = order.service_date or (order.entry_time.date() if order.entry_time else None)
        date_str = _format_date(service_date)
        store_nm = order.store.name if order.store else ""
        local = "Galpão" if order.is_galpon else ""
        dept_label = DEPT_LABELS.get(order.department or "", order.department or "")

        dept_key = (order.department or "").lower()
        if order.external_os_number:
            os_number = order.external_os_number
        elif dept_key in _DEPTS_NO_EXT_OS:
            os_number = ""
        else:
            os_number = order.order_number or f"#{order.id}"

        plate = order.vehicle_plate or ""
        is_courtesy = "Sim" if order.is_courtesy else "Não"
        is_return = "Sim" if order.is_return else "Não"
        consultant_name = (
            order.consultant.name if order.consultant else (order.consultant_name or "")
        )
        vehicle_full = f"{order.vehicle_brand or ''} {order.vehicle_model or ''}".strip()
        color = order.vehicle_color or ""
        notes = _clean_notes(order.notes or "")
        invoice = order.invoice_number or ""

        svc = item.service
        service_code = (svc.code or "") if svc else ""
        service_desc = (svc.name or "") if svc else ""
        tonality = roll.tonality or item.tonality or ""
        roll_code = item.roll_code or ""
        item_value = float(item.unit_price or 0) * int(item.quantity or 1)

        is_verified_mark = "✓" if order.is_verified else ""
        worker_names = ", ".join(w.employee.name for w in (order.workers or []) if w.employee)
        created_at_str = _format_date(order.created_at)
        updated_at_str = _format_date(order.updated_at)
        created_by_name = order.created_by.full_name if order.created_by else ""

        ws.cell(row=row_idx, column=1).value = date_str  # A: Data da O.S
        ws.cell(row=row_idx, column=2).value = store_nm  # B: Loja
        ws.cell(row=row_idx, column=3).value = local  # C: Local
        ws.cell(row=row_idx, column=4).value = dept_label  # D: DPTO
        ws.cell(row=row_idx, column=5).value = plate  # E: Placa/Chassi
        ws.cell(row=row_idx, column=6).value = os_number  # F: O.S
        ws.cell(row=row_idx, column=7).value = is_courtesy  # G: Cortesia?
        ws.cell(row=row_idx, column=8).value = is_return  # H: Retorno?
        ws.cell(row=row_idx, column=9).value = consultant_name  # I: Consultor
        ws.cell(row=row_idx, column=10).value = vehicle_full  # J: Modelo Veículo
        ws.cell(row=row_idx, column=11).value = color  # K: Cor
        ws.cell(row=row_idx, column=12).value = notes  # L: Observações
        ws.cell(row=row_idx, column=13).value = invoice  # M: NFe
        ws.cell(row=row_idx, column=14).value = service_code  # N: Cód Serviço
        ws.cell(row=row_idx, column=15).value = service_desc  # O: Descrição Serviço
        ws.cell(row=row_idx, column=16).value = tonality  # P: Tonalidade
        ws.cell(row=row_idx, column=17).value = roll_code  # Q: Cód Bobina
        ws.cell(row=row_idx, column=18).value = roll_visual_id  # R: ID_Bobina
        ws.cell(row=row_idx, column=19).value = worker_names  # S: Instalador
        vc = ws.cell(row=row_idx, column=20)  # T: Valor
        vc.value = item_value
        vc.number_format = "#,##0.00"
        ws.cell(row=row_idx, column=21).value = is_verified_mark  # U: ✔
        ws.cell(row=row_idx, column=22).value = ""  # V: Obs Conferência
        ws.cell(row=row_idx, column=23).value = created_at_str  # W: Data Cadastro
        ws.cell(row=row_idx, column=24).value = created_by_name  # X: Responsável Cadastro
        ws.cell(row=row_idx, column=25).value = updated_at_str  # Y: Data Alteração
        ws.cell(row=row_idx, column=26).value = ""  # Z: Responsável Alteração
        row_idx += 1

    last_data_row = row_idx - 1
    if last_data_row >= start_row:
        _update_table_ref(ws, last_data_row, "Z")
        subtotal_row = last_data_row + 1
        sc = ws.cell(row=subtotal_row, column=20)  # T: Valor
        sc.value = f"=SUBTOTAL(9,T{start_row}:T{last_data_row})"
        sc.number_format = "#,##0.00"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
